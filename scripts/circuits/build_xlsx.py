#!/usr/bin/env python3
"""
Base circuits — étape 5 : le tableur à vérifier.

Feuilles :
  Synthèse      chiffres clés (formules) et mode d'emploi
  Circuits      une ligne par fiche ; colonnes jaunes = à remplir par Yanis
  Virages       un virage par ligne (circuits au tracé complet)
  Clubs FFSA    circuits, clubs et écoles FFSA de karting, avec la fiche rattachée
  Sources       d'où vient chaque donnée, licences

Usage : python3 build_xlsx.py circuits.json clubs.json sortie.xlsx
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

SURFACES = {"asphalt": "enrobé", "paved": "revêtu", "concrete": "béton", "dirt": "terre", "ground": "terre", "gravel": "gravier"}

F = "Arial"
HEAD = Font(name=F, bold=True, color="FFFFFF", size=10)
HEAD_FILL = PatternFill("solid", fgColor="1F1F1F")
BODY = Font(name=F, size=10)
LINK = Font(name=F, size=10, color="0563C1", underline="single")
TITLE = Font(name=F, bold=True, size=14)
SUB = Font(name=F, size=10, color="666666")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D9D9D9")

KIND = {"epingle": "épingle", "moyen": "moyen", "rapide": "rapide"}
STRUCT = {"circuit": "circuit", "club": "club", "ecole": "école"}


def header(ws, cols, widths, row=1, input_cols=()):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = HEAD
        cell.fill = HEAD_FILL if c not in input_cols else PatternFill("solid", fgColor="BF9000")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=2)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(cols))}{row}"


def main(circ_path, clubs_path, out):
    circuits = json.load(open(circ_path))
    clubs = json.load(open(clubs_path))
    wb = Workbook()

    # ── Circuits ────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Circuits"
    cols = ["Nom", "Commune", "Code postal", "Département", "Région", "Longueur (m)", "Virages",
            "Épingles", "Moyens", "Rapides", "Tracé", "Couvert", "Revêtement", "Largeur (m)",
            "Structures FFSA", "Site web", "Téléphone", "Carte OSM", "Latitude", "Longitude", "Source",
            "Autres tracés sur le site (m)", "Vérifié", "Correction / commentaire"]
    widths = [34, 22, 11, 20, 24, 12, 9, 9, 9, 9, 15, 9, 12, 11, 40, 30, 15, 12, 10, 10, 12, 16, 11, 40]
    header(ws, cols, widths, input_cols=("Vérifié", "Correction / commentaire"))
    for r, c in enumerate(circuits, start=2):
        vals = [c["name"], c["city"], c["postal_code"], c["department"], c["region"], c["length_m"],
                c["corners_count"], c["hairpins_count"], c["medium_count"], c["fast_corners_count"],
                c["trace_status"], {True: "oui", False: "non"}.get(c["indoor"], ""), SURFACES.get(c["surface"], c["surface"]), c["width_m"],
                " · ".join(f"{s['name']} ({STRUCT.get(s['kind'], s['kind'])})" for s in c["ffsa_structures"]),
                c["website"], c["phone"], "voir" if c["osm_url"] else "", c["lat"], c["lon"], c["source"],
                ", ".join(str(x) for x in c.get("annex_tracks_m", [])), None, None]
        for col, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=col, value=v if v != "" else None)
            cell.font = BODY
            cell.border = Border(bottom=THIN)
        if c["osm_url"]:
            cell = ws.cell(row=r, column=18)
            cell.hyperlink = c["osm_url"]
            cell.font = LINK
        if c["website"] and c["website"].startswith("http"):
            ws.cell(row=r, column=16).hyperlink = c["website"]
            ws.cell(row=r, column=16).font = LINK
        for col in (23, 24):
            ws.cell(row=r, column=col).fill = INPUT_FILL
    last = len(circuits) + 1
    dv = DataValidation(type="list", formula1='"OK,À corriger,À supprimer,Doublon"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"W2:W{last}")
    for status, color in (("complet", "E2F0D9"), ("à dessiner", "FCE4D6"), ("partiel", "FFF2CC")):
        ws.conditional_formatting.add(f"K2:K{last}", CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", fgColor=color)))

    # ── Virages ─────────────────────────────────────────────────────────────
    wv = wb.create_sheet("Virages")
    header(wv, ["Circuit", "Commune", "N°", "Nature", "Sens", "Angle (°)", "Rayon mini (m)", "Distance depuis le début du tracé (m)"],
           [34, 22, 6, 10, 8, 10, 14, 22])
    r = 2
    for c in circuits:
        for i, k in enumerate(c["corners"], start=1):
            for col, v in enumerate([c["name"], c["city"], i, KIND[k["kind"]], k["direction"], k["angle_deg"], k["min_radius_m"], k["apex_m"]], start=1):
                wv.cell(row=r, column=col, value=v).font = BODY
            r += 1

    # ── Clubs FFSA ──────────────────────────────────────────────────────────
    wc = wb.create_sheet("Clubs FFSA")
    header(wc, ["Nom", "Type", "Adresse", "Code postal", "Commune", "Fiche circuit rattachée", "Distance à la piste (m)", "Site web", "Téléphone", "Page FFSA"],
           [34, 9, 30, 11, 22, 34, 14, 30, 15, 12])
    for r, c in enumerate(sorted(clubs, key=lambda x: (x["postal_code"], x["name"])), start=2):
        vals = [c["name"], STRUCT.get(c["kind"], c["kind"]), c["street"], c["postal_code"], c["city"], c.get("circuit") or None,
                c.get("track_distance_m"), c["website"] or None, c["phone"] or None, "voir"]
        for col, v in enumerate(vals, start=1):
            wc.cell(row=r, column=col, value=v).font = BODY
        wc.cell(row=r, column=10).hyperlink = c["ffsa_url"]
        wc.cell(row=r, column=10).font = LINK

    # ── Synthèse (formules) ─────────────────────────────────────────────────
    s = wb.create_sheet("Synthèse", 0)
    s.column_dimensions["A"].width = 46
    s.column_dimensions["B"].width = 14
    s.column_dimensions["C"].width = 70
    s["A1"] = "Base des circuits de karting de France"
    s["A1"].font = TITLE
    s["A2"] = "Première version, à vérifier. Les chiffres ci-dessous se recalculent quand tu modifies l'onglet Circuits."
    s["A2"].font = SUB
    rows = [
        ("Fiches circuits", "=COUNTA(Circuits!A2:A{n})"),
        ("…dont tracé complet (longueur et virages calculés)", '=COUNTIF(Circuits!K2:K{n},"complet")'),
        ("…dont tracé partiel dans OpenStreetMap", '=COUNTIF(Circuits!K2:K{n},"partiel")'),
        ("…dont piste encore à dessiner (adresse FFSA seule)", '=COUNTIF(Circuits!K2:K{n},"à dessiner")'),
        ("Fiches liées à au moins une structure FFSA", '=COUNTIF(Circuits!O2:O{n},"?*")'),
        ("Pistes couvertes (indoor)", '=COUNTIF(Circuits!L2:L{n},"oui")'),
        ("Longueur moyenne d'un tracé complet (m)", '=IFERROR(ROUND(AVERAGEIF(Circuits!K2:K{n},"complet",Circuits!F2:F{n}),0),"-")'),
        ("Virages par tracé complet (moyenne)", '=IFERROR(ROUND(AVERAGEIF(Circuits!K2:K{n},"complet",Circuits!G2:G{n}),1),"-")'),
        ("Fiches vérifiées (colonne Vérifié = OK)", '=COUNTIF(Circuits!W2:W{n},"OK")'),
        ("Structures FFSA de karting (circuits, clubs, écoles)", "=COUNTA('Clubs FFSA'!A2:A{m})"),
    ]
    for i, (label, f) in enumerate(rows, start=4):
        s.cell(row=i, column=1, value=label).font = BODY
        cell = s.cell(row=i, column=2, value=f.format(n=last, m=len(clubs) + 1))
        cell.font = Font(name=F, bold=True, size=10)
    base = 4 + len(rows) + 1
    guide = [
        "Comment vérifier",
        "1. Onglet Circuits : les colonnes jaunes sont à toi. « Vérifié » : OK, À corriger, À supprimer ou Doublon. Exemple : SoloKart (Plessé) → OK.",
        "2. « Correction / commentaire » : ce qui est faux (nom, longueur, nombre de virages…), en clair. Exemple : « 1 290 m, 11 virages, ligne de départ côté stand ».",
        "3. Le lien « voir » de la colonne Carte OSM ouvre le tracé : un tracé faux ou incomplet se corrige directement dans OpenStreetMap, et la fiche suivra.",
        "Ce que veulent dire les colonnes",
        "Tracé : complet = boucle fermée de 300 à 2 500 m, longueur et virages calculés ; partiel = la piste existe dans OSM mais en morceaux ; à dessiner = connue de la FFSA, absente d'OSM.",
        "Épingle : virage de 130° ou plus. Rapide : moins de 60°, ou courbe large (rayon ≥ 28 m, ≤ 100°). Moyen : le reste. Calcul sur la géométrie, indépendant du kart et du pilote.",
        "Les chiffres d'un tracé OSM sont des estimations : chaque fichier de session déposé par un pilote sur ce circuit les remplacera par des mesures.",
    ]
    for i, line in enumerate(guide):
        cell = s.cell(row=base + i, column=1, value=line)
        cell.font = Font(name=F, bold=True, size=11) if line in ("Comment vérifier", "Ce que veulent dire les colonnes") else BODY
        s.merge_cells(start_row=base + i, start_column=1, end_row=base + i, end_column=3)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        s.row_dimensions[base + i].height = 30 if len(line) > 90 else 18

    # ── Sources ─────────────────────────────────────────────────────────────
    so = wb.create_sheet("Sources")
    so.column_dimensions["A"].width = 26
    so.column_dimensions["B"].width = 100
    src = [
        ("Source", "Ce qu'on en tire"),
        ("FFSA — Où pratiquer", "Liste publique des circuits, clubs et écoles de karting, avec adresse et coordonnées (ffsa.org/pratiquer/ou-pratiquer). L'adresse est parfois celle du siège et pas celle de la piste."),
        ("OpenStreetMap", "Tracés des pistes (highway=raceway, leisure=track + sport=karting), nom, revêtement, largeur, couvert. © les contributeurs OpenStreetMap, données sous licence ODbL : citer « © OpenStreetMap » partout où un tracé est affiché."),
        ("Nominatim (OSM)", "Recherche des sites de karting par département et par nom, pour retrouver les pistes éloignées de l'adresse FFSA."),
        ("Calcul ApexAI", "Longueur et virages (épingle / moyen / rapide) calculés sur le tracé : scripts/circuits/ du dépôt apex-ai-race-coach."),
        ("Prochaine étape", "Les fichiers de session des pilotes (tracé GPS réel, ligne de départ, virages mesurés) viendront remplacer les estimations OSM."),
    ]
    for i, (a, b) in enumerate(src, start=1):
        so.cell(row=i, column=1, value=a).font = HEAD if i == 1 else Font(name=F, bold=True, size=10)
        so.cell(row=i, column=2, value=b).font = HEAD if i == 1 else BODY
        so.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            so.cell(row=1, column=1).fill = HEAD_FILL
            so.cell(row=1, column=2).fill = HEAD_FILL
        so.row_dimensions[i].height = 32

    wb.save(out)
    print("écrit", out, len(circuits), "fiches")


if __name__ == "__main__":
    main(*sys.argv[1:4])
