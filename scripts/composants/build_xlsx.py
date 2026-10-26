#!/usr/bin/env python3
"""
Base composants : tableur à vérifier.

Usage : python3 build_xlsx.py sortie.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from data import BRAKES, CARBS, CATEGORIES, CHASSIS, ENGINES, MAINTENANCE, SOURCES, TIRES

F = "Arial"
HEAD = Font(name=F, bold=True, color="FFFFFF", size=10)
HEAD_FILL = PatternFill("solid", fgColor="1F1F1F")
INPUT_HEAD = PatternFill("solid", fgColor="BF9000")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
LINK = Font(name=F, size=10, color="0563C1", underline="single")
TITLE = Font(name=F, bold=True, size=14)
SUB = Font(name=F, size=10, color="666666")
WARN = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="D9D9D9")
WRAP = Alignment(vertical="top", wrap_text=True)
INPUTS = ("Vérifié", "Commentaire")

# Écarts entre le catalogue actuel de l'app (table kart_components) et les documents
# (id, catalogue actuel, ce que disent les documents, proposition, source)
ECARTS = [
    ("iame-x30", "Durée de vie 10 h (piston 10 h, réfection 20 h)", "IAME compte en litres de mélange : piston et cage de pied de bielle 100 L en compétition (150 L en loisir), bas moteur 200 L (300 L), bielle 400 L (600 L)", "Compter les litres de mélange consommés (saisie au plein, ou consommation en L/h mesurée une fois par le pilote) et alerter à 100 / 200 / 400 L", "iame_x30"),
    ("iame-x30-junior", "Durée de vie 10 h", "Même moteur que le X30 (bridé) ; aucun manuel Junior séparé trouvé", "Reprendre les intervalles du X30 en litres, à confirmer", "iame_x30"),
    ("rotax-max-125", "Durée de vie 15 h (« ~25 h compétition / 50 h loisir »)", "Démontage et contrôle complet toutes les 50 h par un centre agréé ; contrôles toutes les 2 h et 10 h", "Durée de vie 50 h + rappels 2 h / 10 h (filtre, embrayage, pignons d'équilibrage)", "rotax_evo_2016"),
    ("rotax-junior", "Durée de vie 30 h", "Démontage et contrôle complet toutes les 50 h", "50 h + rappels 2 h / 10 h", "rotax_evo_2016"),
    ("rotax-mini", "Durée de vie 50 h (réfection 100 h)", "Démontage et contrôle complet toutes les 50 h ; aucune valeur à 100 h", "50 h ; retirer la réfection à 100 h", "rotax_evo_2016"),
    ("rotax-micro", "Durée de vie 50 h (réfection 100 h)", "Démontage et contrôle complet toutes les 50 h ; aucune valeur à 100 h", "50 h ; retirer la réfection à 100 h", "rotax_evo_2016"),
    ("rotax-dd2", "Durée de vie 15 h", "Huile de boîte : vidange après les 5 premières heures, puis contrôle toutes les 5 h ; démontage par un centre agréé (intervalle au tableau du manuel, illisible à l'extraction)", "Ajouter le rappel huile de boîte 5 h ; relever l'intervalle de démontage sur le PDF", "rotax_dd2_2024"),
    ("vortex-rok", "Durée de vie 15 h (piston 15 h, réfection 30 h)", "Piston 30 h ; bielle et roulements 60 h (90 h au tableau récapitulatif) ; embrayage 30 h", "Piston 30 h, bas moteur 60 h", "vortex_rokgp"),
    ("vortex-rok-junior", "Durée de vie 15 h", "Le manuel ROK GP couvre aussi les versions Senior et Junior bridées", "Mêmes valeurs que le ROK GP : piston 30 h, bas moteur 60 h", "vortex_rokgp"),
    ("vortex-mini-rok", "Durée de vie 20 h (piston 15 h, réfection 30 h)", "Piston 30 h ; bas moteur 60 h (90 h au tableau)", "Piston 30 h, bas moteur 60 h", "vortex_minirok"),
    ("honda-gx160 / gx200 / gx270 / gx390", "Vidange 20 h, réfection 300 h", "Première vidange à 20 h, puis selon le tableau p. 7 ; aucune « réfection 300 h » dans le manuel", "Garder la première vidange 20 h ; retirer la réfection 300 h", "honda_gx"),
    ("iame-ka100", "Durée de vie 20 h (réfection 40 h)", "Aucun intervalle constructeur trouvé (fiche d'homologation seulement)", "Marquer la valeur comme estimation, à demander à IAME / au motoriste", "ka100_hom"),
    ("tm-kz10c / tm-r1 / tm-r2 / vortex-rvx", "Durée de vie 6 h (piston 4 h, réfection 10 h)", "Aucun manuel constructeur trouvé ; seuls des avis de pilotes et motoristes (3 à 6 h pour le haut moteur)", "Marquer comme estimation", ""),
    ("iame-mini-swift / iame-water-swift / iame-super-x30 / subaru-kx21", "Valeurs d'entretien en heures", "Aucun document constructeur trouvé", "Marquer comme estimation", ""),
    ("komet-k2m / komet-k2h / komet-k1w", "Pneus Komet K2 et K1W", "IAME Eurocup 2026 impose les Komet K3M, K3H et K3W", "Ajouter K3M, K3H, K3W", "eurocup_2026"),
    ("vega-vert-xh3", "Vega XH3", "FFSA Senior / Master / Gentleman 2026 : Vega XH4", "Ajouter XH4", "ffsa_annexe_2026"),
    ("mojo-d2", "MOJO D2", "Rotax Junior MAX : MOJO D2XX", "Ajouter D2XX", "mojo"),
    ("(absents)", "—", "Pneus imposés en France 2026 absents du catalogue : Vega Cadetti, Vega M1, LeCont LOH Option, LeCont LWR, MOJO C2, MOJO CW", "Ajouter", "ffsa_annexe_2026"),
    ("tous les pneus", "Pressions à froid et à chaud « baseline »", "Pas de pression publiée par les fabricants hors MOJO pluie (W5 1,0 bar, CW 0,85 bar en piste détrempée)", "Afficher ces pressions comme point de départ estimé, pas comme une donnée constructeur", "mojo"),
    ("tous les pneus", "Durée de vie en tours (80 à 300)", "Aucune durée de vie publiée", "Marquer comme estimation ; à remplacer par les tours réels des pilotes", ""),
    ("tous les châssis", "Poids 60 kg pour tous", "Aucun poids publié dans les fiches consultées", "Retirer le poids ou le marquer comme estimation", ""),
    ("tonykart-racer-401rr / birel-ry30-s16 / crg-heron", "Tubes seulement", "Empattement 1 045 mm (OTK, Birel), 1 050 mm (CRG Heron) ; axe 50 mm (40 mm en DD2)", "Compléter empattement et axe", "otk_401rr"),
    ("tous les freins", "Durée de vie 600 à 1 000 tours, plaquettes contrôlées tous les 300 tours", "Aucun intervalle publié dans les fiches consultées", "Marquer comme estimation", ""),
    ("carburateurs", "VHSB 34, VHSH 30, HW-27A", "Imposés en France 2026 et absents du catalogue : Tillotson HW-33A (KA100), Dell'Orto PHBG 18 BS (Mini 60)", "Ajouter", "eurocup_2026"),
]


def header(ws, cols, widths, row=1):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = HEAD
        cell.fill = INPUT_HEAD if c in INPUTS else HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=2)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(cols))}{row}"


def rows(ws, data, start=2, src_col=None, link_col=None):
    for r, vals in enumerate(data, start=start):
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v if v not in ("", None) else None)
            cell.font = BODY
            cell.alignment = WRAP
            cell.border = Border(bottom=THIN)
        if link_col:
            url = vals[link_col - 1]
            if url:
                cell = ws.cell(row=r, column=link_col)
                cell.hyperlink = url
                cell.value = "ouvrir"
                cell.font = LINK
    n = start + len(data) - 1
    dv = DataValidation(type="list", formula1='"OK,À corriger,À supprimer"', allow_blank=True)
    ws.add_data_validation(dv)
    return n, dv


def src(key):
    return SOURCES[key][0] if key else "aucune source constructeur trouvée"


def main(out):
    wb = Workbook()

    # ── Synthèse ───────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Synthèse"
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 14
    ws["A1"] = "Base des composants de karting"
    ws["A1"].font = TITLE
    ws["A2"] = "Première version, à vérifier. Chaque valeur vient d'un manuel constructeur ou d'un règlement 2026, avec la page ou l'article."
    ws["A2"].font = SUB
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 28
    counts = [
        ("Catégories françaises 2026 (FFSA, IAME Eurocup, Rotax)", "='Catégories 2026'!A1", f"=COUNTA('Catégories 2026'!B2:B{len(CATEGORIES) + 1})"),
        ("Moteurs avec fiche constructeur", None, f"=COUNTA(Moteurs!B2:B{len(ENGINES) + 1})"),
        ("Opérations d'entretien sourcées", None, f"=COUNTA('Entretien moteurs'!B2:B{len(MAINTENANCE) + 1})"),
        ("Pneus imposés en France 2026", None, f"=COUNTA(Pneus!B2:B{len(TIRES) + 1})"),
        ("Châssis avec fiche constructeur", None, f"=COUNTA(Châssis!B2:B{len(CHASSIS) + 1})"),
        ("Écarts avec le catalogue actuel de l'app", None, f"=COUNTA('Écarts catalogue'!A2:A{len(ECARTS) + 1})"),
    ]
    for i, (label, _, formula) in enumerate(counts, start=4):
        ws.cell(row=i, column=1, value=label).font = BODY
        ws.cell(row=i, column=2, value=formula).font = BOLD
    notes = [
        ("Ce qu'il faut retenir", True),
        ("1. Le catalogue actuel de l'app (75 composants) contient surtout des valeurs estimées : durées de vie, pressions, poids. L'onglet « Écarts catalogue » liste ce qui diffère des documents.", False),
        ("2. IAME compte l'entretien du X30 en litres de mélange, pas en heures (piston 100 L en compétition). Mon Kart compte des heures : il faut ajouter le suivi des litres ou une consommation par heure saisie par le pilote.", False),
        ("3. Rotax prévoit un démontage-contrôle complet toutes les 50 h, avec des contrôles toutes les 2 h et 10 h. L'app alerte aujourd'hui à 15 h sur le 125 MAX evo.", False),
        ("4. L'onglet « Catégories 2026 » donne pour chaque catégorie française le moteur, les pneus, le poids et le carburateur imposés : de quoi préremplir Mon Kart et Réglages à partir de la catégorie choisie.", False),
        ("Comment vérifier", True),
        ("Colonnes jaunes : « Vérifié » (OK, À corriger, À supprimer) et « Commentaire ». Le lien « ouvrir » de l'onglet Sources ouvre le document d'origine.", False),
        ("Rien n'est estimé dans ce fichier : une case vide veut dire que le document ne donne pas la valeur.", False),
    ]
    for i, (t, bold) in enumerate(notes, start=11):
        c = ws.cell(row=i, column=1, value=t)
        c.font = BOLD if bold else BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # ── Catégories 2026 ───────────────────────────────────────────────────
    ws = wb.create_sheet("Catégories 2026")
    cols = ["Organisateur", "Catégorie", "Moteur", "Pneu sec", "Pneu pluie", "Poids mini (kg)", "Carburateur",
            "Transmission", "Châssis", "Autres règles", "Source", "Article", "Vérifié", "Commentaire"]
    header(ws, cols, [16, 22, 40, 22, 18, 10, 24, 18, 40, 40, 36, 14, 11, 30])
    data = [(o, c, m, ps, pp, w, carb, tr, ch, au, src(s), art, None, None) for o, c, m, ps, pp, w, carb, tr, ch, au, s, art in CATEGORIES]
    n, dv = rows(ws, data)
    dv.add(f"M2:M{n}")

    # ── Moteurs ───────────────────────────────────────────────────────────
    ws = wb.create_sheet("Moteurs")
    cols = ["Id catalogue app", "Marque", "Modèle", "Cylindrée (cm³)", "Refroidissement", "Boîte", "Mélange / carburant",
            "Bougie", "Huile de boîte / carter", "Régime", "Autres données", "Source", "Vérifié", "Commentaire"]
    header(ws, cols, [18, 10, 20, 10, 12, 14, 36, 30, 28, 24, 44, 36, 11, 30])
    data = [(i, b, m, cc, cool, gb, mix, plug, oil, rpm, other, src(s), None, None) for i, b, m, cc, cool, gb, mix, plug, oil, rpm, other, s in ENGINES]
    n, dv = rows(ws, data)
    dv.add(f"M2:M{n}")

    # ── Entretien moteurs ─────────────────────────────────────────────────
    ws = wb.create_sheet("Entretien moteurs")
    cols = ["Moteur (id catalogue)", "Élément", "Action", "Intervalle", "Unité", "Usage", "Source", "Page", "Vérifié", "Commentaire"]
    header(ws, cols, [18, 34, 50, 16, 16, 12, 36, 14, 11, 30])
    data = [(e, el, ac, iv, un, us, src(s), pg, None, None) for e, el, ac, iv, un, us, s, pg in MAINTENANCE]
    n, dv = rows(ws, data)
    dv.add(f"I2:I{n}")

    # ── Pneus ─────────────────────────────────────────────────────────────
    ws = wb.create_sheet("Pneus")
    cols = ["Marque", "Modèle", "Type", "Appellation CIK", "Imposé en 2026 par", "Données constructeur",
            "Id catalogue app", "Source", "Vérifié", "Commentaire"]
    header(ws, cols, [10, 22, 8, 12, 44, 50, 18, 36, 11, 30])
    data = [(b, m, t, a, cat, d, i or "absent du catalogue", src(s), None, None) for b, m, t, a, cat, d, i, s in TIRES]
    n, dv = rows(ws, data)
    dv.add(f"I2:I{n}")
    for r in range(2, n + 1):
        if ws.cell(row=r, column=7).value == "absent du catalogue":
            ws.cell(row=r, column=7).fill = WARN

    # ── Châssis ───────────────────────────────────────────────────────────
    ws = wb.create_sheet("Châssis")
    cols = ["Id catalogue app", "Marque", "Modèle", "Tubes Ø (mm)", "Empattement (mm)", "Axe arrière", "Freins",
            "Catégories", "Source", "Vérifié", "Commentaire"]
    header(ws, cols, [20, 12, 30, 10, 14, 30, 50, 22, 36, 11, 30])
    data = [(i or "absent du catalogue", b, m, t, e, a, f, c, src(s), None, None) for i, b, m, t, e, a, f, c, s in CHASSIS]
    n, dv = rows(ws, data)
    dv.add(f"J2:J{n}")

    # ── Freins et carburateurs ────────────────────────────────────────────
    ws = wb.create_sheet("Freins et carburateurs")
    cols = ["Id catalogue app", "Marque", "Modèle", "Montage / usage", "Détail constructeur", "Source", "Vérifié", "Commentaire"]
    header(ws, cols, [18, 12, 24, 30, 60, 36, 11, 30])
    data = [(i or "absent du catalogue", b, m, mo, d, src(s), None, None) for i, b, m, mo, d, s in BRAKES]
    data += [(i or "absent du catalogue", b, m, "carburateur", u, src(s), None, None) for i, b, m, u, s in CARBS]
    n, dv = rows(ws, data)
    dv.add(f"G2:G{n}")

    # ── Écarts catalogue ──────────────────────────────────────────────────
    ws = wb.create_sheet("Écarts catalogue")
    cols = ["Id catalogue app", "Valeur actuelle dans l'app", "Ce que disent les documents", "Proposition", "Source", "Vérifié", "Commentaire"]
    header(ws, cols, [30, 36, 60, 50, 36, 11, 30])
    data = [(i, cur, doc, prop, src(s), None, None) for i, cur, doc, prop, s in ECARTS]
    n, dv = rows(ws, data)
    dv.add(f"F2:F{n}")

    # ── Sources ───────────────────────────────────────────────────────────
    ws = wb.create_sheet("Sources")
    cols = ["Document", "Lien"]
    header(ws, cols, [80, 12])
    rows(ws, [(t, u) for t, u in SOURCES.values()], link_col=2)
    ws.cell(row=len(SOURCES) + 3, column=1, value="Les documents restent la propriété de leurs éditeurs ; ce fichier n'en reprend que des valeurs techniques.").font = SUB

    wb.save(out)
    print("écrit", out)


if __name__ == "__main__":
    main(sys.argv[1])
